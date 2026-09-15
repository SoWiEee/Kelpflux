"""Build the Kelpflux project poster as a single editable PowerPoint slide."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "Kelpflux_poster.pptx"
HERO = ROOT / "assets" / "poster-hero-gpu.png"
TRAINING = ROOT / "assets" / "figures" / "training_convergence.png"

W, H = 85, 110
BG = RGBColor(36, 40, 45)
PANEL = RGBColor(45, 51, 58)
PANEL_DARK = RGBColor(31, 36, 42)
BLUE = RGBColor(91, 143, 190)
BLUE_LIGHT = RGBColor(142, 177, 205)
BLUE_DARK = RGBColor(45, 76, 105)
CREAM = RGBColor(244, 239, 227)
MUTED = RGBColor(183, 187, 190)
DIM = RGBColor(124, 132, 141)
GREEN = RGBColor(132, 171, 145)
ORANGE = RGBColor(195, 157, 105)
FONT = "Noto Sans CJK TC"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.replace("#", "").upper())


def rect(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.8)
    return shape


def textbox(slide, x, y, w, h, text, size=16, color=CREAM, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
            margin=0.12):
    shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Cm(margin)
    tf.margin_right = Cm(margin)
    tf.margin_top = Cm(margin)
    tf.margin_bottom = Cm(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def label(slide, x, y, text, size=12, color=DIM, width=15):
    return textbox(slide, x, y, width, 0.8, text, size=size, color=color,
                   bold=True, margin=0)


def section_title(slide, x, y, title, subtitle=None, width=30):
    textbox(slide, x, y, width, 1.7, title, size=25, color=CREAM, bold=True,
            margin=0)
    rect(slide, x, y + 1.85, 2.8, 0.18, BLUE, BLUE)
    if subtitle:
        textbox(slide, x, y + 2.15, width, 1.35, subtitle, size=12,
                color=MUTED, margin=0)


def circle(slide, x, y, d, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(x), Cm(y), Cm(d), Cm(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def bullet(slide, x, y, text, width, color=CREAM, accent=BLUE, size=15):
    circle(slide, x, y + 0.38, 0.32, accent)
    return textbox(slide, x + 0.65, y, width - 0.65, 1.35, text,
                   size=size, color=color, margin=0)


def metric_panel(slide, x, y, w, title, values, maximum, note):
    rect(slide, x, y, w, 13.2, PANEL_DARK, BLUE_DARK, radius=True)
    textbox(slide, x + 0.9, y + 0.65, w - 1.8, 1.25, title, size=18,
            color=CREAM, bold=True, margin=0)
    textbox(slide, x + 0.9, y + 1.9, w - 1.8, 0.8, note, size=11,
            color=DIM, margin=0)
    row_y = y + 3.05
    bar_x = x + 8.1
    bar_w = w - 14.7
    for method, value, bar_color in values:
        textbox(slide, x + 0.9, row_y - 0.06, 6.7, 0.9, method,
                size=11, color=MUTED, margin=0)
        rect(slide, bar_x, row_y + 0.12, bar_w, 0.42, rgb("#3B434D"), rgb("#3B434D"), radius=True)
        rect(slide, bar_x, row_y + 0.12, max(0.25, bar_w * value / maximum), 0.42,
             bar_color, bar_color, radius=True)
        textbox(slide, x + w - 5.8, row_y - 0.08, 4.8, 0.9,
                f"{value:.1f}", size=12, color=CREAM, bold=True,
                align=PP_ALIGN.RIGHT, margin=0)
        row_y += 1.38
    return rect(slide, x + 0.9, y + 11.65, 0.55, 0.2, BLUE, BLUE)


def pipeline_node(slide, x, y, w, h, number, title, subtitle, accent=BLUE):
    rect(slide, x, y, w, h, PANEL, accent, radius=True)
    circle(slide, x + 0.7, y + 0.55, 1.35, accent)
    textbox(slide, x + 0.7, y + 0.58, 1.35, 1.0, str(number), size=15,
            color=BG, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
            margin=0)
    textbox(slide, x + 2.35, y + 0.52, w - 2.9, 1.15, title, size=16,
            color=CREAM, bold=True, margin=0)
    textbox(slide, x + 2.35, y + 1.85, w - 2.9, h - 2.15, subtitle, size=11,
            color=MUTED, margin=0)


def add_slide():
    prs = Presentation()
    prs.slide_width = Cm(W)
    prs.slide_height = Cm(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

    # Header: title and one generated visual only; the rest stays editable.
    textbox(slide, 3, 2.0, 37, 1.0, "TANET 2026  |  PROJECT POSTER", size=16,
            color=BLUE_LIGHT, bold=True, margin=0)
    textbox(slide, 3, 3.6, 43, 6.0,
            "異質 GPU + MPS\n讓 Slurm 學會選工作、選 GPU",
            size=46, color=CREAM, bold=True, margin=0)
    textbox(slide, 3, 10.2, 40, 3.3,
            "基於 Slurm 與 Kubernetes 架構下 AI 伺服器\nGPU 工作負載智慧排程技術之研究",
            size=20, color=MUTED, margin=0)
    textbox(slide, 3, 14.4, 37, 1.4, "ABC · DEF  ｜ 國立 XX 大學", size=14,
            color=DIM, margin=0)
    slide.shapes.add_picture(str(HERO), Cm(45), Cm(2.0), Cm(37), Cm(15.0))
    rect(slide, 3, 17.9, 79, 0.22, BLUE, BLUE)

    # Problem / gap / objective: one short reading path.
    section_title(slide, 3, 20.0, "為什麼要做？", "問題 → 缺口 → 目標", width=23)
    rect(slide, 3, 23.2, 24.5, 9.2, PANEL_DARK, PANEL_DARK, radius=True)
    label(slide, 4, 24.0, "PROBLEM", color=BLUE_LIGHT, width=12)
    textbox(slide, 4, 25.0, 22.2, 2.6,
            "AI 工作同時提交，\nGPU 變成共享且競爭的資源。",
            size=21, color=CREAM, bold=True, margin=0)
    bullet(slide, 4, 28.0, "GPU 世代不同，效能與記憶體不對稱", 21.8, size=13)
    bullet(slide, 4, 29.4, "MPS 配額會改變共置與完成時間", 21.8, size=13)
    rect(slide, 29.5, 23.2, 24.5, 9.2, PANEL_DARK, PANEL_DARK, radius=True)
    label(slide, 30.5, 24.0, "GAP", color=BLUE_LIGHT, width=12)
    textbox(slide, 30.5, 25.0, 22.0, 2.9,
            "傳統 Slurm 規則\n難以同時看見這些狀態。",
            size=21, color=CREAM, bold=True, margin=0)
    bullet(slide, 30.5, 28.1, "FCFS / Backfill 不學習長期效果", 21.8, size=13)
    bullet(slide, 30.5, 29.5, "既有方法多聚焦單一 GPU 或模擬環境", 21.8, size=13)
    rect(slide, 56, 23.2, 26, 9.2, BLUE_DARK, BLUE_DARK, radius=True)
    label(slide, 57, 24.0, "GOAL", color=CREAM, width=12)
    textbox(slide, 57, 25.0, 23.5, 3.0,
            "在真實 Slurm 路徑中，\n聯合決定工作與 GPU。",
            size=21, color=CREAM, bold=True, margin=0)
    textbox(slide, 57, 28.7, 23.4, 2.7,
            "輸出 GPU placement 與\n25 / 50 / 75 / 100% MPS allocation",
            size=14, color=BLUE_LIGHT, margin=0)

    # Main pipeline, based on assets/architecture.html but simplified for print.
    section_title(slide, 3, 34.3, "一條可落地的排程路徑", "K3s 負責部署與生命週期；Slurm + DRL 負責決策", width=48)
    pipeline_y, node_w, node_h = 38.2, 11.7, 6.1
    nodes = [
        (3, "Job Queue", "工作到達\n前 16 個 pending", BLUE_LIGHT),
        (16.0, "Slurm", "priority\n排序與提交路徑", BLUE),
        (29.0, "DRL Scheduler", "SAC / RDSAC / RLPD\naction mask", BLUE_LIGHT),
        (42.0, "Placement", "選工作 × 選 GPU\nMPS fraction", BLUE),
        (55.0, "GPU Workers", "RTX4070 + RTX3080\n25–100% MPS", BLUE_LIGHT),
        (68.0, "Monitor", "JCT / utilization\nreplay buffer", BLUE),
    ]
    for i, (x, title, subtitle, accent) in enumerate(nodes, 1):
        pipeline_node(slide, x, pipeline_y, node_w, node_h, i, title, subtitle, accent)
        if i < len(nodes):
            arrow(slide, x + node_w, pipeline_y + node_h / 2, x + 12.65,
                  pipeline_y + node_h / 2, color=BLUE_LIGHT, width=1.8)
    textbox(slide, 3, 45.2, 79, 1.55,
            "Fail-safe：決策逾時、不可行或低信心時，回退啟發式，不阻塞 Slurm。",
            size=15, color=CREAM, bold=True, align=PP_ALIGN.CENTER, margin=0)

    # Implementation and evaluation.
    section_title(slide, 3, 48.0, "實作成果與量測結果", "把研究問題落到可部署、可觀察、可比較", width=42)
    # Left implementation panel.
    rect(slide, 3, 51.3, 38.5, 30.1, PANEL_DARK, PANEL_DARK, radius=True)
    label(slide, 4, 52.1, "TRAINING + RUNTIME", color=BLUE_LIGHT, width=20)
    textbox(slide, 4, 53.0, 35.5, 1.6,
            "模擬預訓練 → 真實資料 → RLPD 微調",
            size=19, color=CREAM, bold=True, margin=0)
    slide.shapes.add_picture(str(TRAINING), Cm(4), Cm(55.0), Cm(36.2), Cm(13.0))
    textbox(slide, 4, 68.4, 35.5, 0.8, "SAC / RDSAC 的 reward 與 critic loss 收斂曲線",
            size=10, color=DIM, align=PP_ALIGN.CENTER, margin=0)
    bullet(slide, 4, 70.1, "Slurm submission path plugin：不修改 Slurm core", 35.5, size=13)
    bullet(slide, 4, 72.0, "K3s、DRA、MPS、monitoring 以 Helm 部署", 35.5, size=13)
    bullet(slide, 4, 73.9, "959 次決策：88% RL 主導，12% 低信心回退", 35.5, size=13, accent=GREEN)
    bullet(slide, 4, 75.8, "P99 決策延遲 0.27 ms；最大 7.2 ms，無 150 ms 逾時", 35.5, size=13, accent=GREEN)

    # Right results panels.
    metric_values = [
        ("FCFS", 137.0, DIM), ("Backfill", 136.2, DIM),
        ("Heuristic", 125.2, ORANGE), ("SAC", 126.2, BLUE),
        ("RDSAC-cvar", 130.5, BLUE_LIGHT), ("RLPD", 110.7, GREEN),
    ]
    p99_values = [
        ("FCFS", 255.0, DIM), ("Backfill", 269.8, DIM),
        ("Heuristic", 517.6, ORANGE), ("SAC", 566.5, BLUE),
        ("RDSAC-cvar", 527.7, BLUE_LIGHT), ("RLPD", 509.0, GREEN),
    ]
    metric_panel(slide, 43.0, 51.3, 39.0, "Average JCT（越低越好）", metric_values,
                 150, "mean，依 PDF Table 4；單位：秒")
    metric_panel(slide, 43.0, 65.9, 39.0, "P99 JCT（尾端代價）", p99_values,
                 650, "mean，依 PDF Table 4；單位：秒")
    textbox(slide, 43.9, 79.8, 37.2, 1.25,
            "RLPD 平均 JCT 最低；學習式策略的 P99 約為 FCFS / Backfill 的兩倍。",
            size=12, color=CREAM, bold=True, margin=0)

    # Bottom: testbed, contributions, positioning.
    section_title(slide, 3, 83.2, "研究完成了什麼？", "以實機結果回答研究問題", width=30)
    rect(slide, 3, 86.3, 24.6, 18.6, PANEL_DARK, PANEL_DARK, radius=True)
    label(slide, 4, 87.1, "TESTBED", color=BLUE_LIGHT, width=12)
    textbox(slide, 4, 88.0, 21.8, 1.3, "2-node 異質 GPU", size=20,
            color=CREAM, bold=True, margin=0)
    bullet(slide, 4, 90.0, "RTX 4070｜64 GB RAM", 21.7, size=13)
    bullet(slide, 4, 91.7, "RTX 3080｜8 GB RAM", 21.7, size=13)
    bullet(slide, 4, 93.4, "BERT / ResNet / Qwen / cuBLAS", 21.7, size=13)
    bullet(slide, 4, 95.1, "125 jobs × 8 seeds", 21.7, size=13)
    textbox(slide, 4, 97.7, 21.5, 4.5,
            "K3s 1.34.6\nSlurm 23.11.7\nNVIDIA MPS 25–100%",
            size=13, color=MUTED, margin=0)

    rect(slide, 30.0, 86.3, 25.0, 18.6, BLUE_DARK, BLUE_DARK, radius=True)
    label(slide, 31.0, 87.1, "CONTRIBUTIONS", color=CREAM, width=18)
    bullet(slide, 31.0, 88.6, "聯合 action：job selection × GPU placement", 22.2, size=13, accent=CREAM)
    bullet(slide, 31.0, 91.1, "Fail-safe plugin：逾時即回退啟發式", 22.2, size=13, accent=CREAM)
    bullet(slide, 31.0, 93.6, "真實 Slurm + 異質 GPU + MPS 驗證", 22.2, size=13, accent=CREAM)
    textbox(slide, 31.0, 97.1, 22.2, 4.8,
            "不是只在模擬器裡比較演算法，\n而是把學習決策接到可運作的叢集。",
            size=15, color=CREAM, bold=True, margin=0)

    rect(slide, 57.0, 86.3, 25.0, 18.6, PANEL_DARK, PANEL_DARK, radius=True)
    label(slide, 58.0, 87.1, "POSITIONING", color=BLUE_LIGHT, width=16)
    textbox(slide, 58.0, 88.2, 22.3, 4.1,
            "研究定位\nGPU sharing、placement 與 RL 被放進同一條 Slurm 路徑。",
            size=18, color=CREAM, bold=True, margin=0)
    textbox(slide, 58.0, 93.3, 22.3, 1.0, "[9][10][11][20][21]", size=12,
            color=BLUE_LIGHT, margin=0)
    textbox(slide, 58.0, 95.0, 22.2, 3.8,
            "下一步\n更大叢集 · MIG + MPS · energy-aware reward",
            size=14, color=MUTED, margin=0)

    textbox(slide, 3, 106.2, 79, 1.3,
            "資料與數字取自 assets/TANET.pdf；方法背景：Slurm [1]、NVIDIA MPS [9]、RLPD [20]、Kubernetes DRA [21]。",
            size=9, color=DIM, align=PP_ALIGN.CENTER, margin=0)
    return prs


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    presentation = add_slide()
    presentation.save(OUT)
    print(OUT)
